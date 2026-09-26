import json
import re
import openpyxl

# Load unique catalog
with open('scratch/indiamart_only_catalog.json', 'r', encoding='utf-8') as f:
    unique_catalog = json.load(f)

# Load category listings from crawler and missed categories
all_category_listings = []

def add_listing(cat_name, item):
    title = item.get('title', '')
    img = item.get('image', '')
    price = item.get('price', '')
    pid = str(item.get('id', '')).replace('im-', '')
    
    # find in unique catalog to get clean brand, type, series, model
    matched_f = None
    for u in unique_catalog:
        if u['id'] == pid:
            matched_f = u
            break
            
    if matched_f:
        brand = matched_f['brand']
        ptype = matched_f['type']
        series = matched_f['series']
        model = matched_f['model']
        img_link = matched_f['image_link']
    else:
        brand = ''
        ptype = ''
        series = ''
        model = title
        img_link = img
        
    all_category_listings.append({
        'category': cat_name,
        'brand': brand,
        'type': ptype,
        'series': series,
        'model': model,
        'title': title,
        'price': price,
        'image_link': img_link,
        'product_id': pid
    })

# 1. From all_extracted_category_products.json
try:
    with open('scratch/all_extracted_category_products.json', 'r', encoding='utf-8') as f:
        for it in json.load(f):
            add_listing(it.get('category', 'General'), it)
except Exception as e:
    print("Err cr:", e)

# 2. From missed_categories_products.json
try:
    with open('scratch/missed_categories_products.json', 'r', encoding='utf-8') as f:
        for it in json.load(f):
            add_listing(it.get('category', 'General'), it)
except Exception as e:
    print("Err missed:", e)

print(f"Total Unique Catalog Products: {len(unique_catalog)}")
print(f"Total Category Listing Slots Scraped: {len(all_category_listings)}")

# Write to Excel
wb = openpyxl.Workbook()

# Sheet 1: Unique Products
ws1 = wb.active
ws1.title = "IndiaMART Products"

headers1 = ['Brand', 'type', 'Series', 'model or part number', 'Image link']
font_header = openpyxl.styles.Font(name='Calibri', size=11, bold=True)
font_regular = openpyxl.styles.Font(name='Calibri', size=11)
font_link = openpyxl.styles.Font(name='Calibri', size=11, color="0000FF", underline="single")
align_left = openpyxl.styles.Alignment(horizontal='left', vertical='center')

for c, h in enumerate(headers1, 1):
    cell = ws1.cell(row=1, column=c, value=h)
    cell.font = font_header
    cell.alignment = align_left

for row_idx, prod in enumerate(unique_catalog, 2):
    c1 = ws1.cell(row=row_idx, column=1, value=prod['brand'])
    c1.font = font_regular
    c1.alignment = align_left
    
    c2 = ws1.cell(row=row_idx, column=2, value=prod['type'])
    c2.font = font_regular
    c2.alignment = align_left
    
    c3 = ws1.cell(row=row_idx, column=3, value=prod['series'] if prod['series'] else None)
    c3.font = font_regular
    c3.alignment = align_left
    
    c4 = ws1.cell(row=row_idx, column=4, value=prod['model'])
    c4.font = font_regular
    c4.alignment = align_left
    
    img_val = prod['image_link']
    c5 = ws1.cell(row=row_idx, column=5, value=img_val)
    if img_val and str(img_val).startswith('http'):
        c5.hyperlink = str(img_val)
        c5.font = font_link
    else:
        c5.font = font_regular
    c5.alignment = align_left

ws1.column_dimensions['A'].width = 18
ws1.column_dimensions['B'].width = 25
ws1.column_dimensions['C'].width = 25
ws1.column_dimensions['D'].width = 35
ws1.column_dimensions['E'].width = 50

# Sheet 2: All Category Listings
ws2 = wb.create_sheet(title="All Category Listings (700+)")
headers2 = ['Category Section', 'Brand', 'type', 'Series', 'model or part number', 'Full Product Title', 'Price', 'Image link']

for c, h in enumerate(headers2, 1):
    cell = ws2.cell(row=1, column=c, value=h)
    cell.font = font_header
    cell.alignment = align_left

for row_idx, it in enumerate(all_category_listings, 2):
    ws2.cell(row=row_idx, column=1, value=it['category']).font = font_regular
    ws2.cell(row=row_idx, column=2, value=it['brand']).font = font_regular
    ws2.cell(row=row_idx, column=3, value=it['type']).font = font_regular
    ws2.cell(row=row_idx, column=4, value=it['series']).font = font_regular
    ws2.cell(row=row_idx, column=5, value=it['model']).font = font_regular
    ws2.cell(row=row_idx, column=6, value=it['title']).font = font_regular
    ws2.cell(row=row_idx, column=7, value=it['price']).font = font_regular
    
    c8 = ws2.cell(row=row_idx, column=8, value=it['image_link'])
    if it['image_link'] and str(it['image_link']).startswith('http'):
        c8.hyperlink = str(it['image_link'])
        c8.font = font_link
    else:
        c8.font = font_regular

ws2.column_dimensions['A'].width = 28
ws2.column_dimensions['B'].width = 18
ws2.column_dimensions['C'].width = 22
ws2.column_dimensions['D'].width = 22
ws2.column_dimensions['E'].width = 30
ws2.column_dimensions['F'].width = 45
ws2.column_dimensions['G'].width = 18
ws2.column_dimensions['H'].width = 50

excel_file = 'example sheet (concept automation technologies)_FILLED.xlsx'
wb.save(excel_file)
print(f"Successfully created clean IndiaMART-only Excel sheet: {excel_file}")
