import openpyxl

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

brand_samples = {}
for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    if b and str(b).strip():
        b_name = str(b).strip()
        if b_name not in brand_samples:
            brand_samples[b_name] = []
        if len(brand_samples[b_name]) < 3 and m:
            brand_samples[b_name].append({
                'row': r,
                'brand': b,
                'type': t,
                'series': s,
                'model': m
            })

for b, samples in brand_samples.items():
    print(f"\n--- Brand: {b} ---")
    for s in samples:
        print(f"  Row {s['row']}: Brand='{s['brand']}', type='{s['type']}', Series='{s['series']}', Model='{s['model']}'")
