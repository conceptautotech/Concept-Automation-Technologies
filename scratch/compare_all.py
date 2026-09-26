import json
import openpyxl
import re

# 1. Load Excel models
wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']
excel_models = set()
excel_rows = []
for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    img = ws.cell(row=r, column=5).value
    if m:
        m_str = str(m).strip()
        excel_models.add(m_str.upper())
        excel_rows.append({
            'row': r,
            'brand': str(b).strip() if b else '',
            'type': str(t).strip() if t else '',
            'series': str(s).strip() if s else '',
            'model': m_str,
            'image': str(img).strip() if img else ''
        })

print(f"Total models in Excel sheet: {len(excel_models)} (rows: {len(excel_rows)})")

# 2. Gather all scraped products
catalog_sources = {}

# from catalog.ts
with open('src/data/catalog.ts', 'r', encoding='utf-8', errors='ignore') as f:
    ts_text = f.read()

# Parse JSON objects from catalog.ts
prod_regex = re.finditer(r'{\s*"id":\s*"([^"]+)",\s*"name":\s*"([^"]+)",[\s\S]*?"brand":\s*"([^"]+)",[\s\S]*?"category":\s*"([^"]+)",[\s\S]*?"image":\s*"([^"]+)"', ts_text)
catalog_ts_prods = 0
for m in prod_regex:
    catalog_ts_prods += 1
    pid = m.group(1)
    catalog_sources[pid] = {
        'id': pid,
        'name': m.group(2),
        'brand': m.group(3),
        'category': m.group(4),
        'image': m.group(5),
        'source': 'catalog.ts'
    }
print(f"Parsed {catalog_ts_prods} from catalog.ts. Total unique products: {len(catalog_sources)}")

# from scratch/all_extracted_category_products.json
try:
    with open('scratch/all_extracted_category_products.json', 'r', encoding='utf-8') as f:
        prods = json.load(f)
        for p in prods:
            pid = p.get('id')
            if pid not in catalog_sources:
                catalog_sources[pid] = {
                    'id': pid,
                    'name': p.get('title'),
                    'brand': '',
                    'category': p.get('category'),
                    'image': p.get('image'),
                    'detail_url': p.get('detail_url'),
                    'source': 'crawler'
                }
    print(f"After all_extracted_category_products.json: Total unique products: {len(catalog_sources)}")
except Exception as e:
    print("Error reading all_extracted_category_products.json:", e)

# from scratch/scraped_data.json
try:
    with open('scratch/scraped_data.json', 'r', encoding='utf-8') as f:
        sd = json.load(f)
        print(f"Loaded {len(sd)} items from scraped_data.json")
        for k, v in sd.items():
            if k not in catalog_sources:
                catalog_sources[k] = {
                    'id': k,
                    'name': v.get('description', '').split('\n')[0][:80],
                    'image': v.get('images', [''])[0] if v.get('images') else '',
                    'source': 'scraped_data.json'
                }
    print(f"After scraped_data.json: Total unique products: {len(catalog_sources)}")
except Exception as e:
    print("Error reading scraped_data.json:", e)

# from scratch/product_ids.json
try:
    with open('scratch/product_ids.json', 'r', encoding='utf-8') as f:
        pids = json.load(f)
        print(f"Loaded {len(pids)} items from product_ids.json")
        for item in pids:
            cid = item.get('catalogId') or item.get('indiamartId')
            if cid not in catalog_sources:
                catalog_sources[cid] = {
                    'id': cid,
                    'name': item.get('name'),
                    'detail_url': item.get('url'),
                    'source': 'product_ids.json'
                }
    print(f"After product_ids.json: Total unique products: {len(catalog_sources)}")
except Exception as e:
    print("Error reading product_ids.json:", e)
