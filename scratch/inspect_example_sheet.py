import openpyxl

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    img = ws.cell(row=r, column=5).value
    if m:
        rows.append({
            'row': r,
            'brand': str(b).strip() if b else '',
            'type': str(t).strip() if t else '',
            'series': str(s).strip() if s else '',
            'model': str(m).strip(),
            'image': str(img).strip() if img else ''
        })

print(f'Total valid rows in example sheet: {len(rows)}')
print('\nSample 10 rows:')
for row in rows[:10]:
    print(row)

print('\nSample rows from middle (row 250-260):')
for row in rows[250:260]:
    print(row)

print('\nSample rows from end (last 10):')
for row in rows[-10:]:
    print(row)
