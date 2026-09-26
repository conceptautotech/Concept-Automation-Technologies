import json
import openpyxl
import re

# 1. Load Excel
wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

excel_items = []
excel_models_clean = set()

for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    if m:
        m_str = str(m).strip()
        # normalized model
        norm = re.sub(r'[^A-Z0-9]', '', m_str.upper())
        excel_models_clean.add(norm)
        excel_items.append({
            'row': r,
            'brand': str(b).strip() if b else '',
            'type': str(t).strip() if t else '',
            'series': str(s).strip() if s else '',
            'model': m_str,
            'norm': norm
        })

print(f"Total Excel items: {len(excel_items)}")
print(f"Total Unique Normalized Excel models: {len(excel_models_clean)}")

# 2. Load Master IndiaMART Scraped Products
with open('scratch/master_indiamart_all.json', 'r', encoding='utf-8') as f:
    scraped_data = json.load(f)

print(f"Total Scraped IndiaMART products: {len(scraped_data)}")

# For each scraped product, extract part number / model and check if in Excel
matched = []
unmatched_scraped = []

for pid, item in scraped_data.items():
    title = item.get('title', '')
    specs = item.get('specs', {})
    
    # Candidate part numbers
    candidates = []
    if specs.get('Model Name/Number'):
        candidates.append(specs['Model Name/Number'])
    if specs.get('Model'):
        candidates.append(specs['Model'])
    if specs.get('partNumber'):
        candidates.append(specs['partNumber'])
        
    # Extract from title
    # Matches patterns like FX5U-80MT, 6ES7288-1SR30-0AA0, etc.
    title_matches = re.findall(r'\b([A-Z0-9]{3,}[-\/][A-Z0-9\/-]+|[0-9][A-Z0-9]{5,})\b', title)
    candidates.extend(title_matches)
    
    found = False
    for c in candidates:
        norm_c = re.sub(r'[^A-Z0-9]', '', c.upper())
        if norm_c in excel_models_clean:
            found = True
            break
            
    # Also check if any excel model is contained inside title
    if not found:
        t_norm = re.sub(r'[^A-Z0-9]', '', title.upper())
        for em in excel_models_clean:
            if len(em) >= 5 and em in t_norm:
                found = True
                break
                
    if found:
        matched.append((pid, title))
    else:
        unmatched_scraped.append((pid, item))

print(f"Scraped products already in Excel: {len(matched)}")
print(f"Scraped products PENDING (NOT in Excel): {len(unmatched_scraped)}")

print("\nSample 20 Pending Scraped Products to add to Excel:")
for pid, item in unmatched_scraped[:20]:
    print(f"  [{pid}] {item.get('title')} (Cat: {item.get('category')})")
