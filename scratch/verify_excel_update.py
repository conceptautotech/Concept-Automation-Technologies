import openpyxl

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

print(f"Workbook max_row: {ws.max_row}")

print("\n--- Transition rows 564 to 575 ---")
for r in range(564, 576):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
    hl = ws.cell(row=r, column=5).hyperlink
    hl_str = hl.target[:40] + '...' if hl and hl.target else 'None'
    print(f"Row {r:4d}: {vals[:4]} | Link: {hl_str}")

print("\n--- Rows 1020 to 1032 (End of sheet) ---")
for r in range(1020, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
    hl = ws.cell(row=r, column=5).hyperlink
    hl_str = hl.target[:40] + '...' if hl and hl.target else 'None'
    print(f"Row {r:4d}: {vals[:4]} | Link: {hl_str}")
