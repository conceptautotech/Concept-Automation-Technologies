import openpyxl

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

print("Row 230:", [ws.cell(row=230, column=c).value for c in range(1, 6)])
print("Row 241:", [ws.cell(row=241, column=c).value for c in range(1, 6)])

for r in range(225, 245):
    print(f"Row {r:3d}: {[ws.cell(row=r, column=c).value for c in range(1, 5)]}")
