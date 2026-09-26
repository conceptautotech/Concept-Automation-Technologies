import openpyxl
import re

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

generic_phrases = [
    'touch panel', '& touch panels', 'cylinder sensor', 'proximity sensor', 'magnetic sensor',
    'photoelectric sensor', 'rotary encoder', 'shaft encoder', 'ethernet switch', 'ac drive',
    'vfd', 'plc', 'hmi', 'sensor', 'sensors', 'safety relay', 'allen bradley', 'simatic s7-1200',
    'simatic s7-1500', 's7 1200', 's7-1500', 'ie series', 'ip series', 'general purpose'
]

suspect_rows = []
for r in range(2, ws.max_row + 1):
    m = ws.cell(row=r, column=4).value
    if m:
        m_str = str(m).strip()
        m_lower = m_str.lower()
        if m_lower == 'model' or any(m_lower == g for g in generic_phrases) or any(m_lower.startswith(g) and len(m_str) < len(g) + 5 for g in generic_phrases):
            suspect_rows.append((r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, m_str))

print(f"Total generic / non-model rows: {len(suspect_rows)}")
for r, b, t, m in suspect_rows:
    b_str = str(b) if b else ''
    t_str = str(t) if t else ''
    print(f"  Row {r:4d}: Brand={b_str:<15} Type={t_str:<15} Model='{m}'")
