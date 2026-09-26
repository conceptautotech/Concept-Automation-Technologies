import openpyxl
import re

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

all_rows = []
for r in range(1, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 6)]
    hl = ws.cell(row=r, column=5).hyperlink
    target = hl.target if hl else None
    all_rows.append({
        'orig_row': r,
        'brand': vals[0],
        'type': vals[1],
        'series': vals[2],
        'model': vals[3],
        'image': vals[4],
        'hyperlink': target
    })

header = all_rows[0]
data_rows = all_rows[1:]

generic_models = {
    'MODEL', 'TOUCH PANEL', '& TOUCH PANELS', 'CYLINDER SENSOR', 'PROXIMITY SENSOR',
    'PROXIMITY SENSORS', 'MAGNETIC SENSOR', 'PHOTOELECTRIC SENSOR', 'ROTARY ENCODERS .',
    'SHAFT ENCODER .', 'ETHERNET SWITCH', 'VFD', 'PLC', 'HMI', 'SENSOR', 'SENSORS',
    'SAFETY RELAY', 'ALLEN BRADLEY', 'ALLEN BRADLEY SLC 500', 'SIMATIC S7-1200',
    'SIMATIC S7-1200 BASIC', 'SIMATIC S7-1500', 'S7 1200', 'S7-1500', 'IE SERIES', 'IP SERIES'
}

def normalize_for_dedup(model_str):
    if not model_str: return ''
    m = str(model_str).strip().upper()
    
    # Preserve decimals: 1.5 -> 1DOT5
    m = re.sub(r'(\d+)\.(\d+)', r'\1DOT\2', m)
    
    # Remove leading HMIPFX... -> PFX...
    if m.startswith('HMIPFX'):
        m = m[3:]
    # Remove "- " or "MAKE " or "VFD "
    m = re.sub(r'^[-\s]*(MAKE|VFD|PLC|HMI|AC\s+DRIVES?)\s+', '', m).strip()
    m = re.sub(r'\s+AC\s+DRIVES?$', '', m).strip()
    # Remove trailing -A if FX3U-...-A vs FX3U-...
    if re.match(r'^FX3U-\d+MR/ES-A$', m):
        m = m[:-2]
    # Remove trailing -V1 if NB...-V1
    if re.match(r'^NB\d+W-TW01B-V1$', m):
        m = m[:-3]
    # Remove G120 prefix if G120 CU...
    if m.startswith('G120 '):
        m = m[5:]
    # Standardize alphanumeric
    return re.sub(r'[^A-Z0-9]', '', m)

seen_keys = set()
clean_data_rows = []
removed_rows = []

for r in data_rows:
    m = r['model']
    if m is None or not str(m).strip():
        # Blank row
        continue
        
    m_clean = str(m).strip()
    m_up = m_clean.upper()
    
    # 1. Skip generic non-model placeholders
    if m_up in generic_models:
        removed_rows.append(('GENERIC_PLACEHOLDER', r))
        continue
        
    norm_key = normalize_for_dedup(m_clean)
    if not norm_key or len(norm_key) < 3:
        removed_rows.append(('INVALID_LENGTH', r))
        continue
        
    # Check if this is a truncated duplicate of an existing 16-char Siemens part number
    is_truncated_siemens = False
    if len(norm_key) <= 12 and norm_key.startswith('6S'):
        for sk in seen_keys:
            if len(sk) >= 15 and sk.startswith(norm_key):
                is_truncated_siemens = True
                break
    if is_truncated_siemens:
        removed_rows.append(('TRUNCATED_SIEMENS', r))
        continue
        
    if norm_key in seen_keys:
        removed_rows.append(('DUPLICATE', r))
        continue
        
    seen_keys.add(norm_key)
    # Also standardize the model name if it had leading "- " or "HMIPFX"
    if str(m).strip().startswith('- 22F-A4P2N103'):
        r['model'] = '22F-A4P2N103'
    elif str(m).strip().startswith('HMIPFXGP'):
        r['model'] = str(m).strip()[3:]
    clean_data_rows.append(r)

print(f"Original data rows: {len(data_rows)}")
print(f"Removed rows: {len(removed_rows)}")
print(f"Final clean unique product rows: {len(clean_data_rows)}")

print("\nRemoved items breakdown:")
for reason, item in removed_rows:
    print(f"  [{reason:<20}] Row {item['orig_row']:4d}: Model='{item['model']}' | Brand='{item['brand']}'")
