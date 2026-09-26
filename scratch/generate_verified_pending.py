import json
import re
import openpyxl

# Load clean master products
with open('scratch/clean_master_products.json', 'r', encoding='utf-8') as f:
    master_prods = json.load(f)

# Load existing Excel models
wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

existing_models = set()
for r in range(2, ws.max_row + 1):
    m = ws.cell(row=r, column=4).value
    if m:
        m_str = str(m).strip()
        norm = re.sub(r'[^A-Z0-9]', '', m_str.upper())
        if norm:
            existing_models.add(norm)

# Scraped specs from scraped_data.json
scraped_specs = {}
try:
    with open('scratch/scraped_data.json', 'r', encoding='utf-8') as f:
        sd = json.load(f)
        for k, v in sd.items():
            clean_k = k.replace('im-', '')
            specs_dict = {s['label']: s['value'] for s in v.get('specifications', []) if 'label' in s and 'value' in s}
            scraped_specs[clean_k] = specs_dict
except Exception:
    pass

blacklist = {
    '18DI/12DO', '24DI/16DO', '14DI/10DO', '8DI/6DO', '12DI/8DO', '16DI/16DO', '32DI/32DO',
    'AC/DC/RELAY', 'DC/DC/DC', 'DC/DC/RELAY', 'AC/DC/TRANSISTOR',
    '100-240VAC', '24VDC', '220VAC', '440VAC', '380-480V', '200-240V',
    '150HP', '100HP', '50HP', '10HP', '5HP', '15KW', '11KW', '7.5KW', '5.5KW', '4KW', '2.2KW', '1.5KW', '0.75KW', '0.37KW',
    'MADE IN GERMANY', 'MADE IN JAPAN', 'MADE IN INDIA'
}

def clean_siemens_part_number(s):
    s = s.upper().replace(' ', '')
    if len(s) == 16 and '-' not in s:
        return f"{s[:7]}-{s[7:12]}-{s[12:]}"
    elif len(s) == 18 and '-' not in s:
        return f"{s[:7]}-{s[7:13]}-{s[13:]}"
    return s

def extract_model_part_number(title, specs={}):
    if specs.get('Model Name/Number'):
        m = specs['Model Name/Number'].strip()
        m = re.sub(r'^(Pepperl\s+and\s+Fuchs|Siemens|Mitsubishi|Omron|Allen\s+Bradley|ABB|Schneider|Danfoss|Fuji|Pro-?face|Weintek)\s+', '', m, flags=re.I).strip()
        if len(m) >= 3 and m.upper() not in blacklist and not re.match(r'^\d+DI/\d+DO$', m, re.I):
            return m

    t = title.strip()
    
    # 1. Siemens 6ES7 / 6AV / 6SL / 6ED / 6GK
    siemens_m = re.search(r'\b(6[A-Z0-9]{2,3}[-\s]?[0-9A-Z]{4,5}[-\s]?[0-9A-Z]{3,5})\b', t, re.I)
    if siemens_m:
        val = clean_siemens_part_number(siemens_m.group(1))
        if val not in blacklist:
            return val
            
    siemens_unhyphenated = re.search(r'\b(6[A-Z0-9]{13,17})\b', t, re.I)
    if siemens_unhyphenated:
        return clean_siemens_part_number(siemens_unhyphenated.group(1))

    # 2. Mitsubishi FX3 / FX5 / MR- / FR- / GS / GT / Q
    mitsu_m = re.search(r'\b(FX[0-9][-A-Z0-9_/]+|MR-[-A-Z0-9_]+|FR-[-A-Z0-9_]+|GS[0-9]{4}[-A-Z0-9_]+|GT[0-9]{4}[-A-Z0-9_]+|Q[0-9]{2}[-A-Z0-9_]+)\b', t, re.I)
    if mitsu_m:
        val = mitsu_m.group(1).upper()
        if val not in blacklist and not re.match(r'^\d+DI/\d+DO$', val, re.I):
            return val

    # 3. ABB ACS
    abb_m = re.search(r'\b(ACS[0-9]{3}[-A-Z0-9_/]+)\b', t, re.I)
    if abb_m:
        return abb_m.group(1).upper()

    # 4. Proface PFX
    proface_m = re.search(r'\b(PFX[-A-Z0-9_]+)\b', t, re.I)
    if proface_m:
        return proface_m.group(1).upper()

    # 5. Weintek cMT / MT / TK
    weintek_m = re.search(r'\b(cMT[-0-9A-Z_]+|MT[0-9]{4}[-A-Z0-9_]+|TK[0-9]{4}[-A-Z0-9_]+)\b', t, re.I)
    if weintek_m:
        return weintek_m.group(1)

    # 6. Fuji FRN / MONITOUCH
    fuji_m = re.search(r'\b(FRN[-0-9A-Z_/]+|V[0-9]{3}[-A-Z0-9_]+)\b', t, re.I)
    if fuji_m:
        return fuji_m.group(1).upper()

    # 7. Omron CP / CJ / NX / NB / NA / E2E / E3Z / E6B
    omron_m = re.search(r'\b(CP[0-9][-A-Z0-9_/]+|CJ[0-9][-A-Z0-9_/]+|NX[0-9][-A-Z0-9_/]+|NB[0-9]+[-A-Z0-9_/]+|NA[0-9]+[-A-Z0-9_/]+|E[0-9][-A-Z0-9_/]+)\b', t, re.I)
    if omron_m:
        val = omron_m.group(1).upper()
        if val not in blacklist:
            return val

    # 8. Danfoss 132F or FC
    danfoss_132 = re.search(r'\b(132F\s*[0-9]{4})\b', t, re.I)
    if danfoss_132:
        return danfoss_132.group(1).upper().replace(' ', '')
    danfoss_m = re.search(r'\b(FC[-\s]?[0-9]{2,3}[-A-Z0-9_/]*)\b', t, re.I)
    if danfoss_m:
        return danfoss_m.group(1).upper().replace(' ', '-')

    # 9. Pepperl+Fuchs OBD / NBB / NBN / UB / ML
    pf_m = re.search(r'\b(OBD[-A-Z0-9_/]+|NBB[-A-Z0-9_/]+|NBN[-A-Z0-9_/]+|UB[-A-Z0-9_/]+|ML[0-9]{3}[-A-Z0-9_/]+|ENI58[-A-Z0-9_/]+)\b', t, re.I)
    if pf_m:
        return pf_m.group(1).upper()

    # 10. Phoenix Contact FL SWITCH / QUINT
    phx_m = re.search(r'\b(FL\s+SWITCH\s+[-0-9A-Z_/]+|QUINT[-0-9A-Z_/]+)\b', t, re.I)
    if phx_m:
        return phx_m.group(1).upper()

    # 11. Allen-Bradley: 1769 / 1756 / 2080 / 2711 / 25B
    ab_m = re.search(r'\b(17[0-9]{2}-[-A-Z0-9_/]+|2080-[-A-Z0-9_/]+|2711[-A-Z0-9_/]+|25[A-Z]-[-A-Z0-9_/]+|5069-[-A-Z0-9_/]+)\b', t, re.I)
    if ab_m:
        return ab_m.group(1).upper()

    # 12. Schneider: ATV / TM
    sch_m = re.search(r'\b(ATV[0-9]{3}[-A-Z0-9_/]+|TM[0-9]{3}[-A-Z0-9_/]+)\b', t, re.I)
    if sch_m:
        return sch_m.group(1).upper()

    # Generic code with dash/slash
    gen_m = re.findall(r'\b([A-Z0-9]{3,}[-/][-A-Z0-9_/]+)\b', t)
    for g in gen_m:
        g_up = g.upper()
        if g_up not in blacklist and not re.match(r'^\d+DI/\d+DO$', g_up, re.I) and not re.match(r'^\d+VDC$', g_up, re.I) and len(g_up) >= 4:
            return g_up

    cleaned = re.sub(r'^(Siemens|Mitsubishi|Omron|Allen-?Bradley|ABB|Schneider|Danfoss|Fuji|Pro-?face|Weintek|Autonics|SICK|IFM|Delta|Pilz|Phoenix\s+Contact|Yaskawa|Inovance)\s+', '', t, flags=re.I).strip()
    cleaned = re.sub(r'^(PLC|VFD|HMI|Sensor|Sensors|AC\s+Drive|Drive|Touch\s+Panel|Rotary\s+Encoder|Encoder)\s+', '', cleaned, flags=re.I).strip()
    return cleaned if cleaned else t[:30]

def detect_brand(title, cat):
    comb = f"{title} {cat}".upper()
    if any(k in comb for k in ['SIEMENS', 'SIMATIC', 'SINAMICS', '6ES7', '6AV', '6SL', '6ED', '6GK', 'S7-200', 'S7-1200', 'S7-1500', 'S7-300']): return 'Siemens'
    if any(k in comb for k in ['MITSUBISHI', 'MELSEC', 'FREQROL', 'FX3', 'FX5', 'GOT2000', 'GOT1000']): return 'Mitsubishi'
    if any(k in comb for k in ['OMRON', 'CP1', 'CP2', 'CJ2', 'E2E', 'E3Z', 'SYSMAC']): return 'Omron'
    if any(k in comb for k in ['ALLEN-BRADLEY', 'ALLEN BRADLEY', 'ROCKWELL', 'COMPACTLOGIX', 'MICRO800', '1769', '1756', '2080', 'POINT I/O']): return 'Allen-Bradley'
    if any(k in comb for k in ['ABB', 'ACS560', 'ACS580', 'ACS355', 'ACS880']): return 'ABB'
    if any(k in comb for k in ['SCHNEIDER', 'TELEMECANIQUE', 'ALTIVAR', 'ATV310', 'ATV320', 'ATV630', 'MODICON']): return 'Schneider'
    if any(k in comb for k in ['PRO-FACE', 'PROFACE', 'PFXGP', 'PFXET', 'SP5000']): return 'PROFACE'
    if any(k in comb for k in ['WEINTEK', 'EASYVIEW', 'CMT', 'MT80', 'MT81', 'MT61', 'TK80']): return 'Weintek'
    if any(k in comb for k in ['DANFOSS', 'VLT', 'FC 51', 'FC 302', 'FC-51', 'FC-302', '132F']): return 'Danfoss'
    if any(k in comb for k in ['FUJI', 'FRENIC', 'MONITOUCH', 'FRN']): return 'FUJI'
    if any(k in comb for k in ['PEPPERL+FUCHS', 'PEPPERL AND FUCHS', 'PEPPERL & FUCHS', 'P+F', 'OBD500', 'NBB', 'NBN']): return 'PEPPERL+FUCHS'
    if any(k in comb for k in ['IFM', 'EFFECTOR', 'O5S']): return 'IFM'
    if any(k in comb for k in ['SICK', 'GL6', 'DFS60', 'VTE18']): return 'SICK'
    if any(k in comb for k in ['AUTONICS', 'E50S', 'E40S', 'PR12']): return 'Autonics'
    if any(k in comb for k in ['DELTA', 'DVP', 'DOP', 'MS300']): return 'Delta'
    if any(k in comb for k in ['PILZ', 'PNOZ']): return 'Pilz'
    if any(k in comb for k in ['PHOENIX CONTACT', 'PHOENIX', 'FL SWITCH', 'QUINT']): return 'Phoenix Contact'
    if any(k in comb for k in ['YASKAWA', 'GA700', 'A1000']): return 'Yaskawa'
    if any(k in comb for k in ['INOVANCE', 'MD200', 'MD310', 'MD500']): return 'Inovance'
    if any(k in comb for k in ['HENGSTLER', 'RI58', 'AC58']): return 'Hengstler'
    if any(k in comb for k in ['HEIDENHAIN', 'ERN', 'ECN']): return 'Heidenhain'
    if any(k in comb for k in ['BAUMER', 'EIL580']): return 'Baumer'
    if any(k in comb for k in ['KUEBLER', 'KUBLER', 'SENDIX']): return 'Kuebler'
    if any(k in comb for k in ['TAMAGAWA', 'TS26']): return 'Tamagawa'
    return 'Industrial Automation'

def detect_type(title, cat):
    comb = f"{title} {cat}".upper()
    if 'PHOTOELECTRIC' in comb: return 'PHOTOELECTRIC SENSORS'
    if 'ENCODER' in comb: return 'Encoder'
    if any(k in comb for k in ['PROXIMITY', 'SENSOR', 'ULTRASONIC', 'TEMPERATURE', 'VIBRATION']): return 'SENSORS'
    if any(k in comb for k in ['SERVO DRIVE']): return 'Servo Drive'
    if any(k in comb for k in ['SERVO MOTOR']): return 'SERVO MOTOR'
    if any(k in comb for k in ['VFD', 'AC DRIVE', 'INVERTER', 'VARIABLE FREQUENCY']): return 'VFD'
    if any(k in comb for k in ['HMI', 'TOUCH PANEL', 'TOUCH SCREEN', 'DISPLAY', 'OPERATOR PANEL']): return 'HMI'
    if any(k in comb for k in ['PLC', 'PROGRAMMABLE CONTROLLER', 'CPU', 'MODULE', 'INPUT', 'OUTPUT', 'IO-LINK', 'BASE UNIT']): return 'PLC'
    if any(k in comb for k in ['ETHERNET SWITCH', 'SCALANCE', 'SWITCH']): return 'ETHERNET SWITCH'
    return 'PLC'

def detect_series(title, brand):
    t = title.upper()
    # Mitsubishi
    if 'FX5U' in t or 'MELSEC IQ-F' in t: return 'MELSEC iQ-F'
    if 'FX3U' in t: return 'FX3U'
    if 'FX3G' in t: return 'FX3G'
    if 'FX3S' in t: return 'FX3S'
    if 'GOT2000' in t or 'GS21' in t: return 'GOT2000'
    if 'FR-A800' in t or 'A800' in t: return 'FR-A800'
    if 'FR-D700' in t or 'D700' in t: return 'FR-D700'
    if 'FR-E700' in t or 'E700' in t: return 'FR-E700'
    if 'MR-J4' in t: return 'MELSERVO-J4'
    if 'MR-JE' in t: return 'MELSERVO-JE'
    
    # Siemens
    if 'S7-200 SMART' in t or 'SMART' in t: return 'S7-200 SMART'
    if 'S7-1200' in t: return 'S7-1200'
    if 'S7-1500' in t: return 'S7-1500'
    if 'S7-300' in t: return 'S7-300'
    if 'S7-400' in t: return 'S7-400'
    if 'LOGO' in t: return 'LOGO!'
    if 'ET 200' in t or 'ET200' in t: return 'ET 200'
    if 'COMFORT' in t: return 'Comfort Panel'
    if 'BASIC' in t or 'KTP' in t: return 'Basic Panel'
    if 'V20' in t: return 'SINAMICS V20'
    if 'G120' in t: return 'SINAMICS G120'
    if 'V90' in t: return 'SINAMICS V90'
    if 'SCALANCE' in t: return 'SCALANCE'
    
    # Omron
    if 'CP2E' in t: return 'CP2E'
    if 'CP1E' in t: return 'CP1E'
    if 'CP1L' in t: return 'CP1L'
    if 'CP1H' in t: return 'CP1H'
    if 'CJ2M' in t: return 'CJ2M'
    if 'NB' in t and any(k in t for k in ['NB3', 'NB5', 'NB7', 'NB10']): return 'NB'
    if 'NA5' in t: return 'NA5'
    
    # ABB
    if 'ACS560' in t: return 'ACS560'
    if 'ACS580' in t: return 'ACS580'
    if 'ACS355' in t: return 'ACS355'
    if 'ACS880' in t: return 'ACS880'
    
    # Danfoss
    if 'FC 51' in t or 'FC-51' in t or 'FC51' in t or 'MICRO DRIVE' in t or '132F' in t: return 'VLT Micro Drive FC 51'
    if 'FC 302' in t or 'FC-302' in t or 'FC302' in t: return 'VLT AutomationDrive FC 302'
    
    # Schneider
    if 'ATV310' in t: return 'ATV310'
    if 'ATV320' in t: return 'ATV320'
    if 'ATV630' in t: return 'ATV630'
    
    # Weintek
    if 'CMT' in t: return 'cMT'
    if any(k in t for k in ['MT80', 'MT81', 'MT61', 'TK80']): return 'iE & iP Series'
    
    # Fuji
    if 'MINI' in t or 'C2' in t: return 'FRENIC-Mini (C2)'
    if 'ACE' in t: return 'FRENIC-Ace'
    if 'MEGA' in t: return 'FRENIC-MEGA'
    
    # Proface
    if 'GP4000' in t or 'PFXGP4' in t: return 'GP4000'
    if 'ET6000' in t or 'PFXET6' in t: return 'ET6000'
    if 'SP5000' in t or 'PFXSP' in t: return 'SP5000'
    
    # Phoenix
    if 'FL SWITCH' in t: return 'FL SWITCH'
    if 'QUINT' in t: return 'QUINT POWER'
    
    return ''

final_pending = []
seen_models = set()

for pid, p in master_prods.items():
    title = p.get('title', '')
    cat = p.get('category', '')
    img = p.get('image', '')
    specs = scraped_specs.get(pid, {})
    
    brand = detect_brand(title, cat)
    ptype = detect_type(title, cat)
    series = detect_series(title, brand)
    model = extract_model_part_number(title, specs)
    
    # Format image link
    img_link = img or f"https://www.indiamart.com/proddetail/{pid}.html"
    if 'imimg.com' in img_link:
        img_link = re.sub(r'-\d+x\d+\.', '-1000x1000.', img_link)
        
    norm_m = re.sub(r'[^A-Z0-9]', '', model.upper())
    norm_t = re.sub(r'[^A-Z0-9]', '', title.upper())
    
    if len(norm_m) < 3 or norm_m in existing_models or norm_m in seen_models:
        continue
        
    # Check if exact model in existing excel
    exists = False
    for em in existing_models:
        if len(em) >= 6 and (em == norm_m or em in norm_t):
            exists = True
            break
    if exists:
        continue
        
    seen_models.add(norm_m)
    final_pending.append({
        'brand': brand,
        'type': ptype,
        'series': series,
        'model': model,
        'image_link': img_link,
        'raw_title': title,
        'id': pid
    })

print(f"Total Perfectly Extracted Pending Products: {len(final_pending)}")

# Sort by Brand then Model for clean, professional presentation
final_pending.sort(key=lambda x: (x['brand'], x['type'], x['model']))

with open('scratch/verified_pending_products.json', 'w', encoding='utf-8') as f:
    json.dump(final_pending, f, indent=2)

print("\nSample 25 sorted pending products:")
for p in final_pending[:25]:
    print(f"  {p['brand']:<15} | {p['type']:<15} | {p['series']:<20} | {p['model']}")
