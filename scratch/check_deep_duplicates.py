import openpyxl
import re

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    img = ws.cell(row=r, column=5).value
    
    if m is not None:
        m_str = str(m).strip()
        rows.append({
            'row': r,
            'brand': str(b).strip() if b else '',
            'type': str(t).strip() if t else '',
            'series': str(s).strip() if s else '',
            'model': m_str,
            'image': img
        })

print(f"Total rows inspected: {len(rows)}")

# Check for sub-string duplicates
# e.g., model1 is contained in model2 or vice versa
found_pairs = []
for i in range(len(rows)):
    m1 = rows[i]['model'].upper().replace(' ', '').replace('-', '').replace('/', '')
    if m1 == 'MODEL': continue
    for j in range(i + 1, len(rows)):
        m2 = rows[j]['model'].upper().replace(' ', '').replace('-', '').replace('/', '')
        if m2 == 'MODEL': continue
        if m1 == m2:
            found_pairs.append(('EXACT_NORMALIZED', rows[i], rows[j]))
        elif len(m1) >= 6 and len(m2) >= 6 and (m1 in m2 or m2 in m1) and abs(len(m1) - len(m2)) <= 8:
            found_pairs.append(('SUBSTRING_MATCH', rows[i], rows[j]))

print(f"Found {len(found_pairs)} candidate duplicate pairs:")
for kind, r1, r2 in found_pairs:
    print(f"[{kind}] Row {r1['row']} ('{r1['model']}') vs Row {r2['row']} ('{r2['model']}') | Brands: {r1['brand']} vs {r2['brand']}")
