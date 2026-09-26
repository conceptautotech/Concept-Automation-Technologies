import openpyxl
import re

wb_in = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws_in = wb_in['Sheet1']

# Load our clean_data_rows from clean_all_rows.py logic
generic_models = {
    'MODEL', 'TOUCH PANEL', '& TOUCH PANELS', 'CYLINDER SENSOR', 'PROXIMITY SENSOR',
    'PROXIMITY SENSORS', 'MAGNETIC SENSOR', 'PHOTOELECTRIC SENSOR', 'ROTARY ENCODERS .',
    'SHAFT ENCODER .', 'ETHERNET SWITCH', 'VFD', 'PLC', 'HMI', 'SENSOR', 'SENSORS',
    'SAFETY RELAY', 'ALLEN BRADLEY', 'ALLEN BRADLEY SLC 500', 'SIMATIC S7-1200',
    'SIMATIC S7-1200 BASIC', 'SIMATIC S7-1500', 'S7 1200', 'S7-1500', 'IE SERIES', 'IP SERIES'
}

def normalize_for_dedup(model_str):
    if not model_str: return ''
    m = str(model_str).strip().upper()
    m = re.sub(r'(\d+)\.(\d+)', r'\1DOT\2', m)
    if m.startswith('HMIPFX'):
        m = m[3:]
    m = re.sub(r'^[-\s]*(MAKE|VFD|PLC|HMI|AC\s+DRIVES?)\s+', '', m).strip()
    m = re.sub(r'\s+AC\s+DRIVES?$', '', m).strip()
    if re.match(r'^FX3U-\d+MR/ES-A$', m):
        m = m[:-2]
    if re.match(r'^NB\d+W-TW01B-V1$', m):
        m = m[:-3]
    if m.startswith('G120 '):
        m = m[5:]
    return re.sub(r'[^A-Z0-9]', '', m)

seen_keys = set()
clean_products = []

for r in range(2, ws_in.max_row + 1):
    m = ws_in.cell(row=r, column=4).value
    if m is None or not str(m).strip():
        continue
        
    m_clean = str(m).strip()
    m_up = m_clean.upper()
    
    if m_up in generic_models:
        continue
        
    norm_key = normalize_for_dedup(m_clean)
    if not norm_key or len(norm_key) < 3:
        continue
        
    is_truncated_siemens = False
    if len(norm_key) <= 12 and norm_key.startswith('6S'):
        for sk in seen_keys:
            if len(sk) >= 15 and sk.startswith(norm_key):
                is_truncated_siemens = True
                break
    if is_truncated_siemens:
        continue
        
    if norm_key in seen_keys:
        continue
        
    seen_keys.add(norm_key)
    
    # Clean model if prefix was attached
    final_model = m_clean
    if final_model.startswith('- 22F-A4P2N103'):
        final_model = '22F-A4P2N103'
    elif final_model.startswith('HMIPFXGP'):
        final_model = final_model[3:]
        
    b = ws_in.cell(row=r, column=1).value
    t = ws_in.cell(row=r, column=2).value
    s = ws_in.cell(row=r, column=3).value
    img = ws_in.cell(row=r, column=5).value
    hl = ws_in.cell(row=r, column=5).hyperlink
    hl_target = hl.target if hl else None
    
    clean_products.append({
        'brand': b,
        'type': t,
        'series': s,
        'model': final_model,
        'image': img,
        'hyperlink': hl_target
    })

print(f"Total Unique Clean Products to write: {len(clean_products)}")

# Create a fresh workbook for clean output
wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = 'Sheet1'

# Write Header
headers = ['Brand', 'type', 'Series', 'model or part number', 'Image link']
font_header = openpyxl.styles.Font(name='Calibri', size=11, bold=True)
font_regular = openpyxl.styles.Font(name='Calibri', size=11)
font_link = openpyxl.styles.Font(name='Calibri', size=11, color="0000FF", underline="single")
align_left = openpyxl.styles.Alignment(horizontal='left', vertical='center')

for c, h in enumerate(headers, 1):
    cell = ws_out.cell(row=1, column=c, value=h)
    cell.font = font_header
    cell.alignment = align_left

for row_idx, prod in enumerate(clean_products, 2):
    c_b = ws_out.cell(row=row_idx, column=1, value=prod['brand'])
    c_b.font = font_regular
    c_b.alignment = align_left
    
    c_t = ws_out.cell(row=row_idx, column=2, value=prod['type'])
    c_t.font = font_regular
    c_t.alignment = align_left
    
    c_s = ws_out.cell(row=row_idx, column=3, value=prod['series'])
    c_s.font = font_regular
    c_s.alignment = align_left
    
    c_m = ws_out.cell(row=row_idx, column=4, value=prod['model'])
    c_m.font = font_regular
    c_m.alignment = align_left
    
    img_val = prod['image']
    c_img = ws_out.cell(row=row_idx, column=5, value=img_val)
    if prod['hyperlink']:
        c_img.hyperlink = prod['hyperlink']
        c_img.font = font_link
    elif img_val and str(img_val).startswith('http'):
        c_img.hyperlink = str(img_val)
        c_img.font = font_link
    else:
        c_img.font = font_regular
    c_img.alignment = align_left

# Set reasonable column widths
ws_out.column_dimensions['A'].width = 20
ws_out.column_dimensions['B'].width = 25
ws_out.column_dimensions['C'].width = 30
ws_out.column_dimensions['D'].width = 35
ws_out.column_dimensions['E'].width = 45

excel_file = 'example sheet (concept automation technologies)_FILLED.xlsx'
wb_out.save(excel_file)
print(f"Successfully saved clean deduplicated workbook: {excel_file}")
print(f"Total rows in new file: {ws_out.max_row} (Header + {len(clean_products)} products)")
