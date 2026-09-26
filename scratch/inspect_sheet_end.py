import openpyxl

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

print("Total max_row:", ws.max_row)
for r in range(560, ws.max_row + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
    print(f"Row {r:3d}: {vals}")
