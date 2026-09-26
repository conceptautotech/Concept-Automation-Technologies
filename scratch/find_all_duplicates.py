import openpyxl
import re
from collections import defaultdict

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

rows_by_raw_model = defaultdict(list)
rows_by_norm_model = defaultdict(list)

for r in range(2, ws.max_row + 1):
    b = ws.cell(row=r, column=1).value
    t = ws.cell(row=r, column=2).value
    s = ws.cell(row=r, column=3).value
    m = ws.cell(row=r, column=4).value
    img = ws.cell(row=r, column=5).value
    
    if m and str(m).strip() and str(m).strip().lower() != 'model':
        raw_m = str(m).strip()
        norm_m = re.sub(r'[^A-Z0-9]', '', raw_m.upper())
        
        row_data = {
            'row': r,
            'brand': b,
            'type': t,
            'series': s,
            'model': raw_m,
            'norm': norm_m,
            'image': img
        }
        rows_by_raw_model[raw_m.upper()].append(row_data)
        rows_by_norm_model[norm_m].append(row_data)

print(f"Total rows with model: {sum(len(v) for v in rows_by_raw_model.values())}")

exact_dups = {k: v for k, v in rows_by_raw_model.items() if len(v) > 1}
print(f"Exact case-insensitive model duplicates ({len(exact_dups)}):")
for k, v in exact_dups.items():
    print(f"  '{k}': rows {[item['row'] for item in v]}")

norm_dups = {k: v for k, v in rows_by_norm_model.items() if len(v) > 1 and k not in exact_dups}
print(f"\nNormalized (ignoring dashes/spaces) duplicates ({len(norm_dups)}):")
for k, v in norm_dups.items():
    print(f"  Norm '{k}': {[item['model'] for item in v]} (rows {[item['row'] for item in v]})")
