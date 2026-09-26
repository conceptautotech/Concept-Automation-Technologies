import json
import re
import openpyxl
from build_extractors import detect_brand, detect_type, detect_series, extract_model

wb = openpyxl.load_workbook('example sheet (concept automation technologies)_FILLED.xlsx')
ws = wb['Sheet1']

existing_models = set()
for r in range(2, ws.max_row + 1):
    m = ws.cell(row=r, column=4).value
    if m:
        m_str = str(m).strip()
        norm = re.sub(r'[^A-Z0-9]', '', m_str.upper())
        if norm:
            existing_models.add(norm)

with open('scratch/grand_master_indiamart_products.json', 'r', encoding='utf-8') as f:
    master_prods = json.load(f)

# Also load scraped_specs
scraped_specs = {}
try:
    with open('scratch/scraped_data.json', 'r', encoding='utf-8') as f:
        sd = json.load(f)
        for k, v in sd.items():
            clean_k = k.replace('im-', '')
            specs_dict = {s['label']: s['value'] for s in v.get('specifications', []) if 'label' in s and 'value' in s}
            scraped_specs[clean_k] = specs_dict
except Exception as e:
    pass

pending_to_add = []

for pid, item in master_prods.items():
    title = item.get('title', '')
    cat = item.get('category', '')
    img = item.get('image', '')
    url = item.get('url', f"https://www.indiamart.com/proddetail/{pid}.html")
    specs = scraped_specs.get(pid, {})
    
    brand = detect_brand(title, cat, specs)
    ptype = detect_type(title, cat, specs)
    series = detect_series(title, brand, specs)
    model = extract_model(title, specs)
    
    # Normalize model to check if already in Excel
    norm_m = re.sub(r'[^A-Z0-9]', '', model.upper())
    norm_t = re.sub(r'[^A-Z0-9]', '', title.upper())
    
    already_exists = False
    if norm_m in existing_models:
        already_exists = True
    else:
        for em in existing_models:
            if len(em) >= 6 and (em == norm_m or em in norm_t):
                already_exists = True
                break
                
    if not already_exists:
        # High-res image
        clean_img = img or ''
        if clean_img and 'imimg.com' in clean_img:
            clean_img = re.sub(r'-\d+x\d+\.', '-1000x1000.', clean_img)
            
        pending_to_add.append({
            'id': pid,
            'brand': brand,
            'type': ptype,
            'series': series,
            'model': model,
            'image_link': clean_img or url,
            'title': title
        })

print(f"Total Pending Products Ready to Add: {len(pending_to_add)}")
print("\nSample 30 Pending Products Formatted:")
print(f"{'Brand':<15} | {'Type':<18} | {'Series':<20} | {'Model':<30} | {'Image Link'}")
print("-" * 110)
for p in pending_to_add[:30]:
    img_preview = p['image_link'][:45] + '...' if len(p['image_link']) > 45 else p['image_link']
    print(f"{p['brand']:<15} | {p['type']:<18} | {p['series']:<20} | {p['model']:<30} | {img_preview}")

with open('scratch/pending_products_to_add.json', 'w', encoding='utf-8') as f:
    json.dump(pending_to_add, f, indent=2)
